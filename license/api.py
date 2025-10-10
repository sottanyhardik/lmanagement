# license/api.py
import datetime
import io
from datetime import date

from dateutil.relativedelta import relativedelta
from django.db.models import Q, Sum
from django.http import HttpResponse
from django_filters.rest_framework import DjangoFilterBackend
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from rest_framework import filters, status, permissions, parsers, decorators, response, viewsets
from rest_framework.generics import ListAPIView
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response
from rest_framework.views import APIView

from .filters import LicenseDetailsFilterSet
from .models import (
    LicenseDetailsModel,
    LicenseImportItemsModel,
    LicensePurchase,
    GE, MI, SM, OT, CO, RA, LM,
)
from .serializers import (
    LicenseDetailsSerializer,
    LicenseImportItemsSelectSerializer,
    BiscuitReportSerializer,
    LicensePurchaseSerializer,
)


# ---------- helpers ----------
def _get_bool(param_val):
    if param_val is None:
        return None
    if isinstance(param_val, bool):
        return param_val
    s = str(param_val).strip().lower()
    if s in ("true", "1", "yes", "y"):
        return True
    if s in ("false", "0", "no", "n"):
        return False
    return None


def _get_num(param_val, num_type=float):
    if param_val is None or str(param_val).strip() == "":
        return None
    try:
        return num_type(param_val)
    except (TypeError, ValueError):
        return None


# ---------------- Pagination ----------------
class SmallPagination(PageNumberPagination):
    page_size = 25
    page_size_query_param = "page_size"
    max_page_size = 100


# ---------------- License Details ----------------
class LicenseDetailsViewSet(viewsets.ModelViewSet):
    """
    Default: show NON-EXPIRED (active) unless the client explicitly asks otherwise.
    """
    queryset = (
        LicenseDetailsModel.objects
        .select_related('exporter', 'port')
        .prefetch_related('export_license', 'import_license')
        .all()
        .distinct()
    )
    serializer_class = LicenseDetailsSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_class = LicenseDetailsFilterSet

    search_fields = [
        'license_number', 'file_number', 'notification_number', 'scheme_code',
        'exporter__name', 'port__name',
    ]
    ordering_fields = ['license_date', 'license_expiry_date', 'modified_on']
    ordering = ['-modified_on']

    def get_queryset(self):
        qs = super().get_queryset()
        params = self.request.query_params

        # Balance CIF comparator: ?balance_val=500&balance_cmp=gte|lte (default gte)
        balance_val = _get_num(params.get("balance_val"), float)
        if balance_val is not None:
            cmp_key = (params.get("balance_cmp") or "gte").strip().lower()
            if cmp_key == "lte":
                qs = qs.filter(balance_cif__lte=balance_val)
            else:
                qs = qs.filter(balance_cif__gte=balance_val)

        # Date filters
        from_date = (params.get("from_date") or "").strip()
        to_date = (params.get("to_date") or "").strip()
        if from_date:
            qs = qs.filter(license_date__gte=from_date)
        if to_date:
            qs = qs.filter(license_date__lte=to_date)

        # Text filters
        lic_no = (params.get("license_number") or "").strip()
        if lic_no:
            qs = qs.filter(license_number__icontains=lic_no)

        exporter_in = (params.get("exporter__in") or "").strip()
        if exporter_in:
            ids = [int(x) for x in exporter_in.split(",") if x.isdigit()]
            if ids:
                qs = qs.filter(exporter_id__in=ids)

        port_in = (params.get("port__in") or "").strip()
        if port_in:
            ids = [int(x) for x in port_in.split(",") if x.isdigit()]
            if ids:
                qs = qs.filter(port_id__in=ids)

        # Default active-only unless overridden
        if 'status' not in params and 'is_expired' not in params:
            qs = qs.filter(is_expired=False)

        return qs


# ---------------- Import Items: searchable select ----------------
class LicenseImportItemsSelectView(ListAPIView):
    serializer_class = LicenseImportItemsSelectSerializer
    pagination_class = SmallPagination

    def get_queryset(self):
        p = self.request.query_params
        qs = LicenseImportItemsModel.objects.select_related("license", "hs_code")

        today = date.today()
        one_month_from_now = today + relativedelta(months=1)

        expired_window = _get_bool(p.get("expired"))
        if expired_window is True:
            qs = qs.filter(license__license_expiry_date__lt=one_month_from_now)
        elif expired_window is False:
            qs = qs.filter(license__license_expiry_date__gte=one_month_from_now)

        is_expired_flag = _get_bool(p.get("is_expired"))
        if is_expired_flag is not None:
            qs = qs.filter(license__is_expired=is_expired_flag)

        min_balance_cif = _get_num(p.get("min_balance_cif"), float)
        if min_balance_cif is not None:
            qs = qs.filter(available_value__gte=min_balance_cif)

        min_balance_qty = _get_num(p.get("min_balance_qty"), float)
        if min_balance_qty is not None:
            qs = qs.filter(available_quantity__gte=min_balance_qty)

        sion_id = (p.get("sion_norm_id") or "").strip()
        if sion_id.isdigit():
            qs = qs.filter(license__export_license__norm_class__id=int(sion_id))
        else:
            sion_txt = (p.get("sion_norms") or "").strip()
            if sion_txt:
                qs = qs.filter(license__export_license__norm_class__norm_class__icontains=sion_txt)

        q = (p.get("q") or "").strip()
        if q:
            q_filter = (
                    Q(license__license_number__icontains=q) |
                    Q(description__icontains=q) |
                    Q(hs_code__hs_code__icontains=q)
            )
            if q.isdigit():
                q_filter |= Q(serial_number=int(q))
            qs = qs.filter(q_filter)

        lic_no = (p.get("license_number") or "").strip()
        if lic_no:
            qs = qs.filter(license__license_number__icontains=lic_no)

        desc = (p.get("description") or "").strip()
        if desc:
            qs = qs.filter(description__icontains=desc)

        hs = (p.get("hs_code") or "").strip()
        if hs:
            qs = qs.filter(hs_code__hs_code__icontains=hs)

        notif = (p.get("notification_number") or "").strip()
        if notif:
            qs = qs.filter(license__notification_number__icontains=notif)

        is_null_flag = _get_bool(p.get("is_null"))
        if is_null_flag is not None:
            qs = qs.filter(license__is_null=is_null_flag)

        return qs.distinct().order_by(
            "license__license_expiry_date",
            "license__license_number",
            "serial_number",
        )


# ---------------- Biscuit Report (JSON / XLSX) ----------------
class BiscuitReportAPIView(APIView):
    def get(self, request, party, status_flag, format=None):
        since = datetime.datetime.now() - datetime.timedelta(days=30)
        if 'xlsx' in status_flag:
            format = "xlsx"
            status_flag = status_flag.split(".")[0]
        else:
            format = None
        is_expired = status_flag == "expired"

        qs = LicenseDetailsModel.objects.filter(
            export_license__norm_class__norm_class="E5",
            balance_cif__gte=100,
        )

        party_map = {
            "parle": GE, "mi": MI, "sm": SM, "ot": OT,
            "co": CO, "ra": RA, "lm": LM,
        }

        if is_expired:
            qs = qs.filter(license_expiry_date__lt=since)
        else:
            qs = qs.filter(license_expiry_date__gte=since)

        party_key = (party or "").lower()
        if party_key in party_map:
            qs = qs.filter(purchase_status=party_map[party_key])
            if party_key == "parle":
                qs = qs.filter(exporter__name__icontains="parle")
        else:
            qs = qs.filter(purchase_status=GE).exclude(exporter__name__icontains="parle")

        serializer = BiscuitReportSerializer(qs, many=True)

        wants_xlsx = (format == "xlsx") or (request.query_params.get("format", "").lower() == "xlsx")
        if not wants_xlsx:
            return Response(serializer.data, status=status.HTTP_200_OK)

        return self._as_xlsx_response(serializer.data, party=party_key, status_flag=status_flag)

    def _as_xlsx_response(self, rows, party: str, status_flag: str):
        wb = Workbook()
        ws = wb.active
        ws.title = "Biscuit Report"

        columns = BiscuitReportSerializer.Meta.fields

        header_map = {
            "id": "ID",
            "license_number": "License No",
            "license_expiry_date": "Expiry Date",
            "exporter_name": "Exporter",
            "norm_class": "Norm Class",
            "balance_cif": "Balance CIF",
            "veg_oil_hsn": "Veg Oil HSN",
            "veg_oil_pd": "Veg Oil Description",
            "total_veg_qty": "Total Veg Qty",
            "rbd_qty": "RBD Qty",
            "rbd_cif": "RBD CIF",
            "pko_qty": "PKO Qty",
            "pko_cif": "PKO CIF",
            "veg_qty": "Olive Oil Qty",
            "veg_cif": "Olive Oil CIF",
            "pomace_qty": "Pomace Qty",
            "pomace_cif": "Pomace CIF",
            "ten_restriction": "10% Restriction",
            "juice_hsn": "Juice HSN",
            "juice_pd": "Juice Description",
            "juice_qty": "Juice Qty",
            "juice_cif": "Juice CIF",
            "ff_hsn": "Food Flavour HSN",
            "ff_pd": "Food Flavour Desc",
            "ff_qty": "Food Flavour Qty",
            "df_qty": "Dietary Fibre Qty",
            "f_f_qty": "Fruit Qty",
            "f_f_cif": "Fruit CIF",
            "la_qty": "Leavening Agent Qty",
            "starch_1108": "Starch 1108 Qty",
            "starch__1108_cif": "Starch 1108 CIF",
            "starch_3505": "Starch 3505 Qty",
            "mnm_pd": "Milk/Non-Milk Desc",
            "mnm_qty": "Milk/Non-Milk Qty",
            "cheese_qty": "Cheese Qty",
            "cheese_cif": "Cheese CIF",
            "swp_qty": "SWP Qty",
            "swp_cif": "SWP CIF",
            "wpc_qty": "WPC Qty",
            "wpc_cif": "WPC CIF",
            "pp_hsn": "PP HSN",
            "pp_pd": "PP Description",
            "pp_qty": "PP Qty",
            "get_aluminium": "Aluminium Qty",
            "balance_cif_value": "Balance CIF Value",
        }

        ws.append([header_map.get(col, col) for col in columns])

        for r in rows:
            ws.append([r.get(col, "") for col in columns])

        for idx, col in enumerate(columns, start=1):
            letter = get_column_letter(idx)
            max_len = len(str(header_map.get(col, col)))
            for cell in ws[letter]:
                val = "" if cell.value is None else str(cell.value)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[letter].width = min(max_len + 2, 60)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        today_str = datetime.datetime.now().strftime("%Y%m%d")
        filename = f"biscuit_report_{party}_{status_flag}_{today_str}.xlsx"

        resp = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = f'attachment; filename="{filename}"'
        return resp


# ---------------- License Purchases ----------------
class LicensePurchaseViewSet(viewsets.ModelViewSet):
    """
    CRUD for license purchases/payments with a quick summary endpoint.

    Query params:
      - license=<id>
      - from_date=YYYY-MM-DD
      - to_date=YYYY-MM-DD
      - min_amount=<float>
      - max_amount=<float>
      - search=<text>
    """
    queryset = LicensePurchase.objects.select_related("license").all()
    serializer_class = LicensePurchaseSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = SmallPagination

    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    search_fields = ["reference", "remarks", "license__license_number"]
    ordering_fields = ["purchase_date", "amount_inr", "created_at", "modified_at"]
    ordering = ["-purchase_date"]

    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]

    def get_queryset(self):
        qs = super().get_queryset()
        p = self.request.query_params

        lic = (p.get("license") or "").strip()
        if lic.isdigit():
            qs = qs.filter(license_id=int(lic))

        from_date = (p.get("from_date") or "").strip()
        to_date = (p.get("to_date") or "").strip()
        if from_date:
            qs = qs.filter(purchase_date__gte=from_date)
        if to_date:
            qs = qs.filter(purchase_date__lte=to_date)

        min_amount = _get_num(p.get("min_amount"), float)
        if min_amount is not None:
            qs = qs.filter(amount_inr__gte=min_amount)

        max_amount = _get_num(p.get("max_amount"), float)
        if max_amount is not None:
            qs = qs.filter(amount_inr__lte=max_amount)

        search = (p.get("search") or "").strip()
        if search:
            qs = qs.filter(
                Q(reference__icontains=search) |
                Q(remarks__icontains=search) |
                Q(license__license_number__icontains=search)
            )

        return qs

    @decorators.action(detail=False, methods=["get"], url_path="summary")
    def summary(self, request, *args, **kwargs):
        qs = self.get_queryset()
        total = qs.aggregate(s=Sum("amount_inr"))["s"] or 0
        return response.Response({"total_amount_inr": round(total, 2)})
