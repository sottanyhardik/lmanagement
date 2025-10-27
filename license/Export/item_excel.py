# license/views_xlsx.py
from datetime import date as _date
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from license.models import LicenseDetailsModel
from license.utils import apply_license_filters


class LicenseImportItemsXLSX(View):
    """
    XLSX export with:
      • License fields only on first row per license (includes License No. and FOB INR)
      • Base row per import item (BOE cells blank)
      • Extra rows:
          - Allotments with NO BOE → text in 'Allotments (No BOE)'
          - Each BOE (linked or not) → one row in BOE columns
      • BOE adds: % Premium, Premium per Kg, Premium Amount (Qty*PerKg else %*CIF INR)
      • New columns: FOB INR, Purchase Amount, Net Profit/Loss
      • After each license: two summary rows
          1) Debited totals (CIF FC, CIF INR)
          2) Premium Total and Net Profit/Loss = Premium Total − Purchase Amount
      • NEW: Alternating background color per license block (banded by license)
      • NEW: Port Name and License Date added immediately after License No.
    """

    # ---------------- Queryset ----------------
    def get_queryset(self, request):
        qs = (
            LicenseDetailsModel.objects
            .select_related("exporter", "port")
            .prefetch_related(
                "export_license",
                "import_license",
                "import_license__allotment_details",
                "import_license__allotment_details__allotment",
                "import_license__allotment_details__allotment__company",
                "import_license__allotment_details__allotment__related_company",
                "import_license__allotment_details__allotment__bill_of_entry",
                "import_license__item_details",
                "import_license__item_details__bill_of_entry",
            )
            .order_by("license_expiry_date", "license_number")
        )

        qs = apply_license_filters(qs, request.GET)
        if 'status' not in request.GET and 'is_expired' not in request.GET:
            qs = qs.filter(is_expired=False)

        return qs

    # ---------------- Helpers ----------------
    def _norm_classes_from_export(self, lic):
        values = []
        rel = getattr(lic, "export_license", None)
        for ei in (rel.all() if rel is not None else []):
            val = (
                str(getattr(ei, "norm_class_name", "") or "") or
                str(getattr(getattr(ei, "norm_class", None), "name", "") or getattr(ei, "norm_class", "") or "")
            )
            if val and val not in values:
                values.append(val)
        return ", ".join(values)

    def _sum_fob_inr_from_export(self, lic):
        total = 0.0
        found = False
        rel = getattr(lic, "export_license", None)
        if not rel:
            return None
        candidates = (
            "fob_inr", "fob_value_inr", "fob_amount_inr", "fobInr", "FOB_INR",
            "value_inr", "inr_value",
        )
        for ei in rel.all():
            for fld in candidates:
                if hasattr(ei, fld):
                    v = getattr(ei, fld)
                    if v is not None:
                        try:
                            total += float(v)
                            found = True
                        except Exception:
                            pass
                    break
        return round(total, 2) if found else None

    def _hs_code_of(self, ii):
        for fld in ("hs_code", "hsn_code", "hscode"):
            if hasattr(ii, fld):
                obj = getattr(ii, fld)
                if obj:
                    if hasattr(obj, "code"):
                        return getattr(obj, "code") or getattr(obj, "name", "") or str(obj)
                    if hasattr(obj, "name"):
                        return getattr(obj, "name") or str(obj)
                    return str(obj)
        return ""

    def _desc_of(self, ii):
        for fld in ("description", "item_description", "name"):
            if hasattr(ii, fld) and getattr(ii, fld):
                return str(getattr(ii, fld))
        if hasattr(ii, "item") and ii.item:
            return getattr(ii.item, "name", "") or str(ii.item)
        return ""

    def _qty_of(self, ii):
        for fld in ("quantity", "qty", "import_quantity", "balance_qty"):
            if hasattr(ii, fld):
                v = getattr(ii, fld)
                if v is not None:
                    try:
                        return round(float(v), 2)
                    except Exception:
                        return v
        return None

    def _opening_cif_of(self, lic):
        v = getattr(lic, "opening_balance", None)
        if v is not None:
            try:
                return round(float(v), 2)
            except Exception:
                return v
        return None

    def _purchase_amount_of(self, lic):
        for fld in ("purchase_amount", "purchase_value", "purchase_inr"):
            if hasattr(lic, fld):
                v = getattr(lic, fld)
                if v is not None:
                    try:
                        return round(float(v), 2)
                    except Exception:
                        return v
        return None

    def _extra_wastage_of(self, ai):
        for fld in ("extra_wastage", "wastage_qty", "waste_qty"):
            if hasattr(ai, fld):
                v = getattr(ai, fld)
                if v is not None:
                    try:
                        return round(float(v), 2)
                    except Exception:
                        return v
        return None

    def _allotment_line(self, ai):
        allo = ai.allotment
        qty = ai.qty if ai.qty is not None else ""
        cif_fc = ai.cif_fc if ai.cif_fc is not None else ""
        company = getattr(getattr(allo, "company", None), "name", "") or ""
        extra = self._extra_wastage_of(ai)
        parts = [
            f"Qty {qty:.2f}" if isinstance(qty, (float, int)) else f"Qty {qty}",
            f"CIF.FC {cif_fc:.2f}" if isinstance(cif_fc, (float, int)) else f"CIF.FC {cif_fc}",
            f"Co {company}" if company else "",
            f"Extra {extra:.2f}" if isinstance(extra, (float, int)) else (f"Extra {extra}" if extra else ""),
        ]
        return " | ".join([p for p in parts if p])

    def _boe_company_name(self, boe):
        for attr in ("company", "importer", "exporter", "consignee", "vendor"):
            obj = getattr(boe, attr, None)
            if obj:
                name = getattr(obj, "name", None) or str(obj)
                if name:
                    return name
        for attr in ("company_name", "importer_name", "exporter_name", "consignee_name", "vendor_name"):
            val = getattr(boe, attr, None)
            if val:
                return str(val)
        return ""

    def _build_boe_list_for_item(self, lic, ii):
        boe_map = {}

        rd_list = list(getattr(ii, "item_details").all()) if hasattr(ii, "item_details") else []
        for rd in rd_list:
            boe = getattr(rd, "bill_of_entry", None)
            if not boe:
                continue
            bid = getattr(boe, "id", None)
            comp = self._boe_company_name(boe)
            if bid not in boe_map:
                boe_map[bid] = {
                    "no": getattr(boe, "bill_of_entry_number", "") or "",
                    "date": getattr(boe, "bill_of_entry_date", None),
                    "company": comp,
                    "qty": 0.0,
                    "cif_fc": 0.0,
                    "cif_inr": 0.0,
                }
            for attr, key in (("qty", "qty"), ("cif_fc", "cif_fc"), ("cif_inr", "cif_inr")):
                try:
                    boe_map[bid][key] += float(getattr(rd, attr, 0) or 0)
                except Exception:
                    pass
            if not boe_map[bid]["company"] and comp:
                boe_map[bid]["company"] = comp

        allotment_items = list(getattr(ii, "allotment_details").all()) if hasattr(ii, "allotment_details") else []
        for ai in allotment_items:
            allo = ai.allotment
            if hasattr(allo, "bill_of_entry"):
                for boe in allo.bill_of_entry.all():
                    bid = getattr(boe, "id", None)
                    comp = self._boe_company_name(boe)
                    if bid not in boe_map:
                        boe_map[bid] = {
                            "no": getattr(boe, "bill_of_entry_number", "") or "",
                            "date": getattr(boe, "bill_of_entry_date", None),
                            "company": comp,
                            "qty": None,
                            "cif_fc": None,
                            "cif_inr": None,
                        }
                    elif not boe_map[bid]["company"] and comp:
                        boe_map[bid]["company"] = comp

        def _sort_key(v):
            d = v["date"]
            return (d or _date.min, v["no"])

        return sorted(boe_map.values(), key=_sort_key)

    # ---------------- Build XLSX ----------------
    def get(self, request, *args, **kwargs):
        wb = Workbook()
        ws = wb.active
        ws.title = "License Import Items"

        # Column map (updated: added port and license_date after lic_no)
        COLS = {
            "lic_no": 1, "port": 2, "license_date": 3, "expiry": 4, "exporter": 5, "norms": 6,
            "open_cif": 7, "bal_cif": 8, "fob_inr": 9, "purchase": 10,
            "hs": 11, "desc": 12, "qty": 13, "allot": 14,
            "boe_no": 15, "boe_date": 16, "boe_comp": 17, "boe_qty": 18,
            "boe_cif_fc": 19, "boe_cif_inr": 20, "boe_pct": 21, "boe_perkg": 22,
            "boe_prem": 23, "net_pl": 24,
        }

        header = [
            "License No.",
            "Port Name",
            "License Date",
            "License Expiry Date",
            "License Exporter",
            "License Norms",
            "License Opening CIF",
            "License Balance CIF",
            "FOB INR",
            "Purchase Amount",
            "HS Code",
            "Description",
            "Quantity",
            "Allotments (No BOE)",
            "BOE No.",
            "BOE Date",
            "BOE Company",
            "BOE Qty",
            "BOE CIF FC",
            "BOE CIF INR",
            "% Premium",
            "Premium per Kg",
            "Premium Amount",
            "Net Profit/Loss",
        ]

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True)
        body_font = Font(name="Calibri", size=10)
        fill_header = PatternFill("solid", fgColor="F2F2F2")
        fill_total = PatternFill("solid", fgColor="FFF4CC")  # totals highlight

        # License band fills (alternate per license)
        fill_band_a = PatternFill("solid", fgColor="F7FBFF")  # very light blue
        fill_band_b = PatternFill("solid", fgColor="FDF7FF")  # very light lilac

        thin = Side(style="thin", color="BFBFBF")
        border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

        al_left = Alignment(vertical="top", wrap_text=True)
        al_right = Alignment(vertical="top", horizontal="right", wrap_text=True)
        al_center = Alignment(vertical="top", horizontal="center", wrap_text=True)

        # Header
        ws.append(header)
        for col in range(1, len(header) + 1):
            c = ws.cell(row=1, column=col)
            c.font = header_font
            c.fill = fill_header
            c.border = border_all
            c.alignment = al_center

        def style_row(r, band_fill, numeric_cols=()):
            """Apply fonts/borders/alignment and band fill to row r."""
            for col_idx in range(1, len(header) + 1):
                cell = ws.cell(row=r, column=col_idx)
                cell.font = body_font
                cell.border = border_all
                # number formats for numeric columns (apply even if formula string is present)
                if col_idx in numeric_cols:
                    cell.alignment = al_right
                    if col_idx == COLS["boe_pct"]:
                        cell.number_format = "0.00%"
                    else:
                        # default numeric format (works for both numeric value & formula result)
                        cell.number_format = "0.00"
                else:
                    cell.alignment = al_left
                # apply band fill
                cell.fill = band_fill

        row_idx = 2
        qs = self.get_queryset(request)
        band_toggle = False  # flips for each license

        for lic in qs:
            # choose band color for this license
            band_toggle = not band_toggle
            band_fill = fill_band_a if band_toggle else fill_band_b

            norms = self._norm_classes_from_export(lic)
            bal_cif_attr = getattr(lic, "get_balance_cif", None)
            if callable(bal_cif_attr):
                try:
                    bal_cif_val = bal_cif_attr()
                    bal_cif_val = round(float(bal_cif_val), 2) if bal_cif_val is not None else None
                except Exception:
                    bal_cif_val = None
            else:
                try:
                    bal_cif_val = round(float(bal_cif_attr), 2) if bal_cif_attr is not None else None
                except Exception:
                    bal_cif_val = bal_cif_attr

            open_cif_val = self._opening_cif_of(lic)
            purchase_val = self._purchase_amount_of(lic)
            fob_inr_val = self._sum_fob_inr_from_export(lic)
            exporter_name = getattr(getattr(lic, "exporter", None), "name", "") or ""
            license_number = getattr(lic, "license_number", "") or ""
            port_name = getattr(getattr(lic, "port", None), "name", "") or ""
            license_date_val = getattr(lic, "license_date", None)

            items_qs = getattr(lic, "import_license", None)
            items = list(items_qs.all()) if items_qs is not None else []
            license_first_row = None
            boe_rows_for_license = []

            for idx, ii in enumerate(items):
                # base row
                expiry_str = lic.license_expiry_date.strftime("%d-%b-%Y") if getattr(lic, "license_expiry_date",
                                                                                     None) and idx == 0 else ""
                license_date_str = license_date_val.strftime("%d-%b-%Y") if license_date_val and idx == 0 else ""
                row = [""] * len(header)
                row[COLS["lic_no"] - 1] = license_number if idx == 0 else ""
                row[COLS["port"] - 1] = port_name if idx == 0 else ""
                row[COLS["license_date"] - 1] = license_date_str if idx == 0 else ""
                row[COLS["expiry"] - 1] = expiry_str
                row[COLS["exporter"] - 1] = exporter_name if idx == 0 else ""
                row[COLS["norms"] - 1] = norms if idx == 0 else ""
                row[COLS["open_cif"] - 1] = open_cif_val if idx == 0 else None
                row[COLS["bal_cif"] - 1] = bal_cif_val if idx == 0 else None
                row[COLS["fob_inr"] - 1] = fob_inr_val if idx == 0 else None
                row[COLS["purchase"] - 1] = purchase_val if idx == 0 else None
                row[COLS["hs"] - 1] = self._hs_code_of(ii)
                row[COLS["desc"] - 1] = self._desc_of(ii)
                row[COLS["qty"] - 1] = self._qty_of(ii)

                ws.append(row)
                style_row(
                    row_idx,
                    band_fill,
                    numeric_cols=(
                        COLS["open_cif"], COLS["bal_cif"], COLS["fob_inr"], COLS["purchase"],
                        COLS["qty"], COLS["boe_qty"], COLS["boe_cif_fc"], COLS["boe_cif_inr"],
                        COLS["boe_prem"], COLS["boe_pct"], COLS["boe_perkg"]
                    ),
                )
                if idx == 0:
                    license_first_row = row_idx
                row_idx += 1

                # allotments with NO BOE
                allotment_items = list(getattr(ii, "allotment_details").all()) if hasattr(ii, "allotment_details") else []
                for ai in allotment_items:
                    allo = ai.allotment
                    has_boe = hasattr(allo, "bill_of_entry") and getattr(allo.bill_of_entry, "exists", lambda: False)()
                    if not has_boe:
                        line = self._allotment_line(ai)
                        ws.append([""] * (COLS["allot"] - 1) + [line] + [""] * (len(header) - COLS["allot"]))
                        style_row(
                            row_idx,
                            band_fill,
                            numeric_cols=(COLS["open_cif"], COLS["bal_cif"], COLS["fob_inr"], COLS["purchase"],
                                          COLS["qty"])
                        )
                        row_idx += 1

                # BOE rows
                boe_list = self._build_boe_list_for_item(lic, ii)
                for v in boe_list:
                    date_str = v["date"].strftime("%d-%b-%Y") if v["date"] else ""
                    # create a full row starting at BOE columns
                    boe_row = [""] * (COLS["boe_no"] - 1) + [
                        v["no"],
                        date_str,
                        v["company"],
                        v["qty"],
                        v["cif_fc"],
                        v["cif_inr"],
                        0.0,  # % Premium (as decimal)
                        0.0,  # Premium per Kg
                        None,  # Premium Amount (formula below)
                        "",  # Net P/L (summary rows only)
                    ]
                    ws.append(boe_row)
                    # number formats & band fill
                    style_row(
                        row_idx,
                        band_fill,
                        numeric_cols=(COLS["boe_qty"], COLS["boe_cif_fc"], COLS["boe_cif_inr"], COLS["boe_pct"],
                                      COLS["boe_perkg"], COLS["boe_prem"]),
                    )
                    # Premium Amount = IF(PerKg>0, Qty*PerKg, IF(%>0, %*CIF_INR, 0))
                    r = row_idx
                    qty_cell = f"{get_column_letter(COLS['boe_qty'])}{r}"
                    perkg_cell = f"{get_column_letter(COLS['boe_perkg'])}{r}"
                    pct_cell = f"{get_column_letter(COLS['boe_pct'])}{r}"
                    cif_inr_cell = f"{get_column_letter(COLS['boe_cif_inr'])}{r}"
                    prem_cell = ws.cell(row=r, column=COLS["boe_prem"])
                    prem_cell.value = (
                        f"=IF({perkg_cell}>0,{qty_cell}*{perkg_cell},IF({pct_cell}>0,{cif_inr_cell}*{pct_cell},0))"
                    )
                    prem_cell.number_format = "0.00"
                    boe_rows_for_license.append(r)
                    row_idx += 1

            # ---- Summary rows (kept in yellow, not banded) ----
            if boe_rows_for_license:
                r1, rN = boe_rows_for_license[0], boe_rows_for_license[-1]
                col_fc = get_column_letter(COLS["boe_cif_fc"])
                col_inr = get_column_letter(COLS["boe_cif_inr"])
                col_prem = get_column_letter(COLS["boe_prem"])

                # Debited totals
                ws.append([""] * len(header))
                r_tot1 = row_idx
                ws.cell(row=r_tot1, column=COLS["lic_no"]).value = "Debited Totals"
                ws.cell(row=r_tot1, column=COLS["boe_cif_fc"]).value = f"=SUM({col_fc}{r1}:{col_fc}{rN})"
                ws.cell(row=r_tot1, column=COLS["boe_cif_inr"]).value = f"=SUM({col_inr}{r1}:{col_inr}{rN})"
                for c in range(1, len(header) + 1):
                    cell = ws.cell(row=r_tot1, column=c)
                    cell.font = Font(name="Calibri", size=10, bold=True)
                    cell.fill = fill_total
                    cell.border = border_all
                    cell.alignment = al_right if c in (COLS["boe_cif_fc"], COLS["boe_cif_inr"]) else al_left
                    if c in (COLS["boe_cif_fc"], COLS["boe_cif_inr"]):
                        cell.number_format = "0.00"
                row_idx += 1

                # Premium total + Net P/L
                ws.append([""] * len(header))
                r_tot2 = row_idx
                ws.cell(row=r_tot2, column=COLS["lic_no"]).value = "Premium & Net"
                ws.cell(row=r_tot2, column=COLS["boe_prem"]).value = f"=SUM({col_prem}{r1}:{col_prem}{rN})"
                prem_total_cell = f"{get_column_letter(COLS['boe_prem'])}{r_tot2}"
                if license_first_row:
                    purch_cell = f"{get_column_letter(COLS['purchase'])}{license_first_row}"
                else:
                    purch_cell = "0"
                ws.cell(row=r_tot2, column=COLS["net_pl"]).value = f"=IFERROR({prem_total_cell}-{purch_cell},\"\")"
                for c in range(1, len(header) + 1):
                    cell = ws.cell(row=r_tot2, column=c)
                    cell.font = Font(name="Calibri", size=10, bold=True)
                    cell.fill = fill_total
                    cell.border = border_all
                    if c in (COLS["boe_prem"], COLS["net_pl"], COLS["purchase"]):
                        cell.alignment = al_right
                        cell.number_format = "0.00"
                    else:
                        cell.alignment = al_left
                row_idx += 1

        # Column widths (updated indices)
        widths = {
            1: 22,  # License No.
            2: 20,  # Port Name
            3: 16,  # License Date
            4: 18,  # Expiry
            5: 28,  # Exporter
            6: 28,  # Norms
            7: 16,  # Opening CIF
            8: 16,  # Balance CIF
            9: 16,  # FOB INR
            10: 18, # Purchase Amount
            11: 14, # HS Code
            12: 40, # Description
            13: 12, # Quantity
            14: 48, # Allotments
            15: 18, # BOE No.
            16: 14, # BOE Date
            17: 32, # BOE Company
            18: 12, # BOE Qty
            19: 14, # BOE CIF FC
            20: 14, # BOE CIF INR
            21: 12, # % Premium
            22: 14, # Premium per Kg
            23: 16, # Premium Amount
            24: 18, # Net P/L
        }
        for idx, w in widths.items():
            ws.column_dimensions[get_column_letter(idx)].width = w

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}1"

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        resp = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = (
            f'inline; filename="LicenseImportItems_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        )
        return resp
